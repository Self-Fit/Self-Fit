"""Read the supplied workbook without changing it; retain every source row."""
import datetime
import hashlib
import json
import os
import pathlib
import shutil
import sys
import openpyxl

source = pathlib.Path(sys.argv[1]).resolve()
destination = pathlib.Path(sys.argv[2]).resolve()
destination.mkdir(parents=True, exist_ok=True, mode=0o700)
os.chmod(destination, 0o700)
digest = hashlib.sha256(source.read_bytes()).hexdigest()
backup = destination / "source.xlsx"
if backup.exists() and hashlib.sha256(backup.read_bytes()).hexdigest() != digest:
    raise RuntimeError("A different source workbook already exists in this import directory")
book = openpyxl.load_workbook(source, read_only=True, data_only=True)
def serialize(value):
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.isoformat()
    raise TypeError(type(value).__name__)
data = {"filename": source.name, "sha256": digest, "sheets": {},
        "dateEpoch": book.epoch.isoformat(), "formulaCells": {}}
for sheet in book:
    data["sheets"][sheet.title] = [list(row) for row in sheet.iter_rows(values_only=True)]
book.close()
formulas = openpyxl.load_workbook(source, read_only=True, data_only=False)
for sheet in formulas:
    cached_rows = data["sheets"][sheet.title]
    cells = []
    for row in sheet.iter_rows():
        for cell in row:
            if cell.data_type == "f":
                cells.append({"cell": cell.coordinate, "formula": cell.value,
                              "cached": cached_rows[cell.row - 1][cell.column - 1]})
    data["formulaCells"][sheet.title] = cells
formulas.close()
out = destination / "source.json"
out.write_text(json.dumps(data, ensure_ascii=False, default=serialize), encoding="utf-8")
os.chmod(out, 0o600)
shutil.copyfile(source, backup)
os.chmod(backup, 0o600)
print(json.dumps({"sha256": digest, "sheets": {k:len(v) for k,v in data["sheets"].items()}, "output": str(out)}, ensure_ascii=False))
